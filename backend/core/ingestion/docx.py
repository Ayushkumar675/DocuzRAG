"""DOCX paragraph and table extraction."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from core.ingestion.parser import ParsedItem
from core.ingestion.tables import extract_docx_tables


def parse_docx(path: Path, document_id: str, filename: str) -> list[ParsedItem]:
    """Extract paragraphs and tables with document-local references."""

    document = Document(path)
    items: list[ParsedItem] = []
    section: str | None = None

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            section = text
        items.append(
            ParsedItem(
                content_type="TEXT",
                text=text,
                page_number=None,
                source_reference=f"document:{document_id}#paragraph={index}",
                section=section,
                metadata={"format": "docx", "paragraph_index": index, "style": style_name},
            )
        )

    for index, table in enumerate(extract_docx_tables(document)):
        items.append(
            ParsedItem(
                content_type="TABLE",
                text=table.text,
                page_number=None,
                source_reference=f"document:{document_id}#table={index}",
                section=section,
                metadata={"format": "docx", **table.metadata},
            )
        )
    return items
